import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

from app.schemas.scorecard import Scorecard, ScorecardRow

SUFFICIENCY_LABEL = {
    "confirmed": "Подтверждено",
    "partial": "Частично",
    "not_demonstrated": "Не раскрыто",
    "no_data": "Данных нет",
}
SUFFICIENCY_COLOR = {
    "confirmed": "D1FAE5",
    "partial": "FEF3C7",
    "not_demonstrated": "E2E8F0",
    "no_data": "E0F2FE",
}
PROVENANCE_LABEL = {
    "independent": "самостоятельно",
    "clarified": "после уточнения",
    "corrected": "после исправления",
    "hinted": "после подсказки",
    "unknown": "—",
}
EVIDENCE_LABEL = {
    "demonstrated": "продемонстрировано",
    "told": "со слов кандидата",
    "mixed": "смешанно",
    "none": "—",
}
_ACCENT = RGBColor(0x4F, 0x46, 0xE5)
_MUTED = RGBColor(0x64, 0x74, 0x8B)


def _shade(cell, color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


def _counts(scorecard: Scorecard) -> str:
    total = len(scorecard.rows)
    confirmed = sum(1 for r in scorecard.rows if r.sufficiency == "confirmed")
    gaps = sum(1 for r in scorecard.rows if r.sufficiency == "no_data")
    disputed = sum(1 for r in scorecard.rows if r.disputed)
    return (
        f"Компетенций: {total} · подтверждено: {confirmed} · "
        f"пробелов: {gaps} · спорных выводов: {disputed}"
    )


def scorecard_to_markdown(scorecard: Scorecard) -> str:
    lines = [f"# Карточка кандидата — {scorecard.role} {scorecard.level}".strip(), ""]
    lines.append(f"_{_counts(scorecard)}_")
    lines.append("")
    lines.append("## Внутренняя карточка")
    lines.append("")
    for row in scorecard.rows:
        lines.append(f"### {row.name} — {SUFFICIENCY_LABEL.get(row.sufficiency)}")
        lines.append(f"- Ход решения: {PROVENANCE_LABEL.get(row.provenance)}")
        lines.append(f"- Тип основания: {EVIDENCE_LABEL.get(row.evidence_kind)}")
        if row.observation:
            lines.append(f"- Наблюдение: {row.observation}")
        if row.conclusion:
            lines.append(f"- Вывод: {row.conclusion}")
        if row.disputed:
            lines.append(f"- ⚠ Спорно: {row.dispute_reason}")
        for question in row.open_questions:
            lines.append(f"- Открытый вопрос: {question}")
        lines.append("")

    lines.append("## Вопросы для следующего этапа")
    lines.append("")
    for item in scorecard.next_questions:
        lines.append(f"- {item.question} ({item.why_unknown})")
    lines.append("")

    lines.append("## Версия для кандидата")
    lines.append("")
    lines.append("**Сильные стороны:** " + ", ".join(scorecard.candidate_view.strengths))
    lines.append("**Зоны роста:** " + ", ".join(scorecard.candidate_view.growth_areas))
    return "\n".join(lines)


def _set_base_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)


def _row_block(document: Document, row: ScorecardRow) -> None:
    heading = document.add_paragraph()
    run = heading.add_run(row.name)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = _ACCENT

    meta = document.add_paragraph()
    tag = meta.add_run(
        f"{SUFFICIENCY_LABEL.get(row.sufficiency)}  ·  "
        f"{PROVENANCE_LABEL.get(row.provenance)}  ·  "
        f"{EVIDENCE_LABEL.get(row.evidence_kind)}"
    )
    tag.font.size = Pt(9)
    tag.font.color.rgb = _MUTED

    if row.observation:
        document.add_paragraph(row.observation)
    if row.conclusion:
        p = document.add_paragraph()
        p.add_run("Вывод: ").bold = True
        p.add_run(row.conclusion)
    if row.disputed:
        warn = document.add_paragraph()
        wr = warn.add_run(f"⚠ Спорный вывод: {row.dispute_reason}")
        wr.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)
    for question in row.open_questions:
        document.add_paragraph(question, style="List Bullet")


def scorecard_to_docx(scorecard: Scorecard) -> bytes:
    document = Document()
    _set_base_font(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    trun = title.add_run("Карточка кандидата")
    trun.bold = True
    trun.font.size = Pt(22)
    subtitle = document.add_paragraph()
    srun = subtitle.add_run(f"{scorecard.role} · {scorecard.level}".strip(" ·"))
    srun.font.size = Pt(12)
    srun.font.color.rgb = _MUTED
    summary = document.add_paragraph()
    sr = summary.add_run(_counts(scorecard))
    sr.font.size = Pt(10)
    sr.font.color.rgb = _ACCENT

    document.add_heading("Сводка по компетенциям", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    for cell, text in zip(header, ("Компетенция", "Достаточность", "Ход решения")):
        cell.paragraphs[0].add_run(text).bold = True
        _shade(cell, "4F46E5")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in scorecard.rows:
        cells = table.add_row().cells
        cells[0].text = row.name
        cells[1].text = SUFFICIENCY_LABEL.get(row.sufficiency, "")
        _shade(cells[1], SUFFICIENCY_COLOR.get(row.sufficiency, "FFFFFF"))
        cells[2].text = PROVENANCE_LABEL.get(row.provenance, "")

    document.add_heading("Внутренняя карточка", level=1)
    for row in scorecard.rows:
        _row_block(document, row)

    if scorecard.next_questions:
        document.add_heading("Вопросы для следующего этапа", level=1)
        for item in scorecard.next_questions:
            p = document.add_paragraph(style="List Bullet")
            p.add_run(item.question).bold = True
            if item.why_unknown:
                note = p.add_run(f"  — {item.why_unknown}")
                note.font.color.rgb = _MUTED

    document.add_heading("Версия для кандидата", level=1)
    document.add_paragraph("Сильные стороны", style="Heading 3")
    for strength in scorecard.candidate_view.strengths or ["—"]:
        document.add_paragraph(strength, style="List Bullet")
    document.add_paragraph("Зоны роста", style="Heading 3")
    for area in scorecard.candidate_view.growth_areas or ["—"]:
        document.add_paragraph(area, style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
